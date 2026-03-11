"""Service for generating ICP Excel templates and parsing uploaded ICP Excel files.

Single-ICP format: one sheet with Field/Value rows.
"""

from io import BytesIO
from typing import List
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="5C2D8F", end_color="5C2D8F", fill_type="solid")
EXAMPLE_FONT = Font(name="Calibri", italic=True, color="888888", size=10)
FIELD_FONT = Font(name="Calibri", bold=True, size=11)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
FIELD_ALIGNMENT = Alignment(horizontal="left", vertical="center")
VALUE_ALIGNMENT = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="D6D6D6"),
    right=Side(style="thin", color="D6D6D6"),
    top=Side(style="thin", color="D6D6D6"),
    bottom=Side(style="thin", color="D6D6D6"),
)

# Template field definitions: (field_label, example_value)
# These match the frontend ICPDefinition type structure exactly.
TEMPLATE_FIELDS = [
    ("ICP Name", "Enterprise SaaS - US"),
    ("Description", "Target mid-market SaaS companies in the US for our analytics platform"),
    ("Industries", "Technology / SaaS, Technology / Data & Analytics"),
    ("Countries", "United States, Canada"),
    ("Priority Areas", "San Francisco Bay Area, New York, Austin"),
    ("Employees Min", "200"),
    ("Employees Max", "5000"),
    ("Revenue Min", "20000000"),
    ("Revenue Max", "500000000"),
    ("Revenue Currency", "USD"),
    ("Low Cost Center", "No"),
    ("Target Offerings", "Cloud analytics platform, Data pipeline automation"),
    ("Offerings Condition", "OR"),
    ("Urgency Signals", "Recent funding round, Leadership change, Regulatory deadline"),
    ("Urgency Condition", "OR"),
    ("Budget Signals", "Recent fundraise >$10M, IT budget expansion, New CTO hire"),
    ("Budget Condition", "OR"),
    ("Target Roles", "CTO, VP Engineering, Head of Data"),
]


def generate_icp_template() -> BytesIO:
    """Generate a single-ICP Excel template with one sheet (field/value rows)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "ICP Configuration"

    # Header row
    ws.cell(row=1, column=1, value="Field")
    ws.cell(row=1, column=2, value="Value")
    for col in (1, 2):
        cell = ws.cell(row=1, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER

    # Data rows
    for i, (field, example) in enumerate(TEMPLATE_FIELDS, start=2):
        field_cell = ws.cell(row=i, column=1, value=field)
        field_cell.font = FIELD_FONT
        field_cell.alignment = FIELD_ALIGNMENT
        field_cell.border = THIN_BORDER

        value_cell = ws.cell(row=i, column=2, value=example)
        value_cell.font = EXAMPLE_FONT
        value_cell.alignment = VALUE_ALIGNMENT
        value_cell.border = THIN_BORDER

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 65

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _split_csv(value: str) -> list[str]:
    """Split a comma-separated string into a list of stripped, non-empty strings."""
    return [v.strip() for v in value.split(",") if v.strip()]


def _parse_industries(value: str) -> list[dict]:
    """Parse industries from 'Vertical / Sub-vertical, Vertical / Sub-vertical' format."""
    entries = _split_csv(value)
    industries = []
    for entry in entries:
        parts = entry.split("/", 1)
        vertical = parts[0].strip()
        sub_vertical = parts[1].strip() if len(parts) > 1 else None
        if vertical:
            industries.append({"vertical": vertical, "sub_vertical": sub_vertical})
    return industries


def _safe_int(value, default: int = 0) -> int:
    """Safely convert a value to int."""
    if value is None:
        return default
    try:
        return int(float(str(value).strip().replace(",", "")))
    except (ValueError, TypeError):
        return default


def parse_icp_excel(buffer: BytesIO) -> List[dict]:
    """Parse an uploaded Excel file (single-ICP format) and return ICP configuration.

    Returns a single-element list:
    [{"name": str, "description": str, "config": dict, "warnings": list[str]}]
    """
    wb = load_workbook(buffer, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    warnings = []

    # Build field map: field_label (lowercase) -> value
    field_map: dict[str, str] = {}
    for row in ws.iter_rows(min_row=2, max_col=2, values_only=True):
        field = row[0] if row and row[0] else None
        value = row[1] if row and len(row) > 1 and row[1] else None
        if field:
            key = str(field).strip().lower()
            field_map[key] = str(value).strip() if value is not None else ""

    wb.close()

    # Extract name and description
    name = field_map.get("icp name", "")
    description = field_map.get("description", "")

    if not name:
        warnings.append("No ICP name found — please fill in the 'ICP Name' field")

    # Parse low_cost_center boolean
    lcc_raw = field_map.get("low cost center", "").lower()
    low_cost_center = lcc_raw in ("yes", "true", "1")

    # Parse condition fields (default to OR)
    offerings_condition = field_map.get("offerings condition", "OR").upper()
    if offerings_condition not in ("AND", "OR"):
        offerings_condition = "OR"
    urgency_condition = field_map.get("urgency condition", "OR").upper()
    if urgency_condition not in ("AND", "OR"):
        urgency_condition = "OR"
    budget_condition = field_map.get("budget condition", "OR").upper()
    if budget_condition not in ("AND", "OR"):
        budget_condition = "OR"

    # Build config matching frontend ICPDefinition structure
    config = {
        "firmographic_details": {
            "industry_types": _parse_industries(field_map.get("industries", "")),
            "geography": {
                "countries": _split_csv(field_map.get("countries", "")),
                "priority_areas": _split_csv(field_map.get("priority areas", "")),
            },
            "revenue_range": {
                "min": _safe_int(field_map.get("revenue min"), 10000000),
                "max": _safe_int(field_map.get("revenue max"), 500000000),
                "currency": field_map.get("revenue currency", "USD").upper() or "USD",
            },
            "employee_range": {
                "min": _safe_int(field_map.get("employees min"), 50),
                "max": _safe_int(field_map.get("employees max"), 1500),
            },
            "low_cost_center": low_cost_center,
        },
        "target_capability": {
            "offerings": _split_csv(field_map.get("target offerings", "")),
            "condition": offerings_condition,
        },
        "urgency_signals": {
            "signals": _split_csv(field_map.get("urgency signals", "")),
            "condition": urgency_condition,
        },
        "budget_signals": {
            "signals": _split_csv(field_map.get("budget signals", "")),
            "condition": budget_condition,
        },
        "authority_roles": {
            "target_roles": _split_csv(field_map.get("target roles", "")),
        },
    }

    if not config["target_capability"]["offerings"]:
        warnings.append("No target offerings found")
    if not config["firmographic_details"]["geography"]["countries"]:
        warnings.append("No target regions/countries found")

    return [{
        "name": name,
        "description": description,
        "config": config,
        "warnings": warnings,
    }]
