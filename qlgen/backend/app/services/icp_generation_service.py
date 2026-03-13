import asyncio
import csv
import io
import json
import logging
from datetime import datetime

import boto3

from app.config import get_settings

logger = logging.getLogger(__name__)
settings = get_settings()

_bedrock_client = None

MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB
MAX_EXTRACTED_TEXT = 50_000  # chars


def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    """Extract text content from uploaded file (PDF, DOCX, XLSX, TXT, CSV).

    Raises ValueError on unsupported type, empty content, or oversized file.
    """
    if len(file_bytes) > MAX_FILE_SIZE:
        raise ValueError("File exceeds 10 MB limit")

    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext == "txt":
        text = file_bytes.decode("utf-8", errors="replace")
    elif ext == "csv":
        decoded = file_bytes.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(decoded))
        text = "\n".join(", ".join(row) for row in reader)
    elif ext == "pdf":
        try:
            import pdfplumber
        except ImportError:
            raise ValueError("PDF support requires pdfplumber (pip install pdfplumber)")
        pages_text = []
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    pages_text.append(page_text)
        text = "\n\n".join(pages_text)
    elif ext == "docx":
        try:
            from docx import Document
        except ImportError:
            raise ValueError("DOCX support requires python-docx (pip install python-docx)")
        doc = Document(io.BytesIO(file_bytes))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    elif ext == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        rows_text = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cell_values = [str(c) for c in row if c is not None]
                if cell_values:
                    rows_text.append(", ".join(cell_values))
        wb.close()
        text = "\n".join(rows_text)
    else:
        raise ValueError(f"Unsupported file type: .{ext}. Supported: PDF, DOCX, XLSX, TXT, CSV")

    text = text.strip()
    if not text:
        raise ValueError("Uploaded file contains no extractable text")

    if len(text) > MAX_EXTRACTED_TEXT:
        text = text[:MAX_EXTRACTED_TEXT]

    return text


ICP_CONFIG_KEYS = [
    "firmographic_details",
    "target_capability",
    "urgency_signals",
    "budget_signals",
    "authority_roles",
]

ICP_GENERATION_SYSTEM_PROMPT = """You are an expert B2B sales strategist. Given a natural language description of an ideal customer, generate a structured ICP (Ideal Customer Profile) configuration as a JSON object.

IMPORTANT: The current year is {current_year}. Use this year in the ICP name, not any other year.

Return ONLY a JSON object with exactly these keys:
- "name": A short descriptive name for this ICP (e.g., "MidMarket US ECommerce {current_year}")
- "description": A one-sentence summary of the ICP
- "config": An object with exactly these 5 keys:

1. "firmographic_details": {
     "industry_types": [{"vertical": string, "sub_vertical": string or null}],
     "geography": {"countries": [strings]},
     "revenue_range": {"min": int, "max": int, "currency": "USD"|"EUR"|"GBP"|"INR"},
     "employee_range": {"min": int, "max": int},
     "low_cost_center": boolean
   }
2. "target_capability": {"offerings": [strings], "condition": "OR"}
3. "urgency_signals": {"signals": [strings], "condition": "OR"}
4. "budget_signals": {"signals": [strings], "condition": "OR"}
5. "authority_roles": {"target_roles": [strings]}

Rules:
- All array fields must have at least 1 item
- Use sensible defaults: "mid-market" = 100-2000 employees, "enterprise" = 2000+, "SMB" = 10-200
- Revenue should be in whole numbers (e.g., 10000000 for $10M)
- If the user doesn't specify a region, default to "United States"
- If the user doesn't specify a currency, default to "USD"
- low_cost_center should be false unless specifically mentioned
- condition should be "OR" unless the description implies all signals must match
- Generate 3-5 items per array field based on the description
- Return ONLY the JSON object, no markdown fences, no explanation"""


def _get_bedrock_client():
    global _bedrock_client
    if _bedrock_client is None:
        _bedrock_client = boto3.client(
            "bedrock-runtime",
            region_name=settings.AWS_REGION,
        )
    return _bedrock_client


def _extract_json_from_response(text: str) -> dict:
    """Extract JSON from response text, stripping markdown fences if present."""
    text = text.strip()

    # Strip markdown code fences
    if text.startswith("```"):
        first_newline = text.find("\n")
        if first_newline != -1:
            text = text[first_newline + 1:]
        closing = text.rfind("```")
        if closing != -1:
            text = text[:closing]
        text = text.strip()

    # Find outermost JSON object by brace-depth matching
    if "{" not in text:
        raise ValueError("No JSON object found in response")

    start = text.index("{")
    depth = 0
    for i, char in enumerate(text[start:], start):
        if char == "{":
            depth += 1
        elif char == "}":
            depth -= 1
            if depth == 0:
                return json.loads(text[start:i + 1])

    # If we didn't find matching braces, try parsing from start of {
    return json.loads(text[start:])


def _validate_icp_structure(parsed: dict) -> dict:
    """Validate and normalize the parsed ICP structure."""
    # If Claude returned the config directly without name/description wrapper
    if "config" not in parsed and any(k in parsed for k in ICP_CONFIG_KEYS):
        parsed = {
            "name": "AI Generated ICP",
            "description": None,
            "config": parsed,
        }

    if "config" not in parsed:
        raise ValueError("Response missing 'config' key")

    config = parsed["config"]

    # Ensure all 5 required keys exist
    for key in ICP_CONFIG_KEYS:
        if key not in config:
            raise ValueError(f"Config missing required key: {key}")

    # Coerce numeric fields in firmographic_details
    fd = config.get("firmographic_details", {})
    rev = fd.get("revenue_range", {})
    for field in ["min", "max"]:
        if field in rev:
            try:
                rev[field] = int(rev[field])
            except (ValueError, TypeError):
                pass
    emp = fd.get("employee_range", {})
    for field in ["min", "max"]:
        if field in emp:
            try:
                emp[field] = int(emp[field])
            except (ValueError, TypeError):
                pass

    return parsed


def generate_icp_config(description: str, file_text: str | None = None) -> dict:
    """Call Bedrock to generate an ICP config from a natural language description."""
    client = _get_bedrock_client()

    # Build user message combining file content and description
    parts = []
    if file_text:
        parts.append(f"Here is the content of an uploaded document:\n\n{file_text}\n")
    if description.strip():
        parts.append(f"Additional description:\n{description}")
    elif not file_text:
        parts.append(description)

    user_message = "\n\n".join(parts) if parts else description

    # Inject current year into the system prompt
    current_year = datetime.now().year
    system_prompt = ICP_GENERATION_SYSTEM_PROMPT.replace("{current_year}", str(current_year))

    body = json.dumps({
        "anthropic_version": "bedrock-2023-05-31",
        "system": system_prompt,
        "messages": [
            {"role": "user", "content": user_message}
        ],
        "max_tokens": 4096,
        "temperature": 0.3,
    })

    response = client.invoke_model(
        modelId=settings.BEDROCK_MODEL_ID,
        body=body,
        contentType="application/json",
        accept="application/json",
    )

    result = json.loads(response["body"].read())
    text = result["content"][0]["text"]

    parsed = _extract_json_from_response(text)
    validated = _validate_icp_structure(parsed)

    return validated


async def generate_icp_config_async(description: str, file_text: str | None = None) -> dict:
    """Async wrapper around generate_icp_config."""
    return await asyncio.to_thread(generate_icp_config, description, file_text)
