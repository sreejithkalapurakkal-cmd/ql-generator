import csv
from io import BytesIO, StringIO
from uuid import UUID

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy import or_, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.company import Company
from app.models.company_stage import CompanyStageResult


def _extract_signal_text(company: Company, stage_name: str) -> str:
    """Extract semicolon-separated signal names from a company's stage result evidence."""
    if not hasattr(company, 'stage_results') or not company.stage_results:
        return ""
    for sr in company.stage_results:
        if sr.stage == stage_name and sr.evidence:
            signals = []
            for item in sr.evidence:
                if isinstance(item, dict):
                    name = item.get("signal_name") or item.get("signal") or ""
                    if name:
                        signals.append(str(name))
            return "; ".join(signals)
    return ""


async def _fetch_companies(run_id: UUID, db: AsyncSession, scope: str = "all"):
    stmt = (
        select(Company)
        .where(Company.pipeline_run_id == run_id)
        .options(selectinload(Company.contacts), selectinload(Company.stage_results))
        .order_by(Company.name)
    )
    if scope == "final":
        # Exclude disqualified companies; for multi-step runs also
        # require promoted=True (non-promoted are early-stage discards).
        stmt = stmt.where(Company.qualification != "disqualified")
        stmt = stmt.where(
            or_(
                Company.promoted == True,   # noqa: E712  — multi-step: explicitly promoted
                Company.promoted.is_(None),  # single-step: promoted is NULL for all
            )
        )
    result = await db.execute(stmt)
    return result.scalars().unique().all()


async def generate_xlsx(run_id: UUID, db: AsyncSession, scope: str = "all") -> BytesIO:
    companies = await _fetch_companies(run_id, db, scope=scope)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Qualified Leads"

    headers = [
        "Serial#", "Company Name", "Website", "Geo/City",
        "Contact Name", "Designation", "LinkedIn", "Email",
        "Phone", "Final Score", "Budget Score", "Urgency Score",
        "Deal Hotness", "Hotness Tier", "Budget Signal", "Urgency Signal",
    ]

    header_fill = PatternFill(start_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    row_num = 2
    serial = 1
    for company in companies:
        score_display = company.final_score if company.final_score is not None else "N/A"
        city_str = ", ".join(filter(None, [company.city, company.state_region, company.country]))
        budget_display = company.budget_signal_score if company.budget_signal_score is not None else ""
        urgency_display = company.urgency_signal_score if company.urgency_signal_score is not None else ""
        hotness_display = company.deal_hotness_score if company.deal_hotness_score is not None else ""
        tier_display = company.deal_hotness_tier or ""
        budget_signal_text = _extract_signal_text(company, "budget_signals")
        urgency_signal_text = _extract_signal_text(company, "urgency_signals")

        if company.contacts:
            for contact in company.contacts:
                ws.cell(row=row_num, column=1, value=serial)
                ws.cell(row=row_num, column=2, value=company.name)
                ws.cell(row=row_num, column=3, value=company.website)
                ws.cell(row=row_num, column=4, value=city_str)
                ws.cell(row=row_num, column=5, value=contact.full_name)
                ws.cell(row=row_num, column=6, value=contact.designation)
                ws.cell(row=row_num, column=7, value=contact.linkedin_url)
                ws.cell(row=row_num, column=8, value=contact.email)
                ws.cell(row=row_num, column=9, value=contact.phone)
                ws.cell(row=row_num, column=10, value=score_display)
                ws.cell(row=row_num, column=11, value=budget_display)
                ws.cell(row=row_num, column=12, value=urgency_display)
                ws.cell(row=row_num, column=13, value=hotness_display)
                ws.cell(row=row_num, column=14, value=tier_display)
                ws.cell(row=row_num, column=15, value=budget_signal_text)
                ws.cell(row=row_num, column=16, value=urgency_signal_text)
                serial += 1
                row_num += 1
        else:
            ws.cell(row=row_num, column=1, value=serial)
            ws.cell(row=row_num, column=2, value=company.name)
            ws.cell(row=row_num, column=3, value=company.website)
            ws.cell(row=row_num, column=4, value=city_str)
            ws.cell(row=row_num, column=10, value=score_display)
            ws.cell(row=row_num, column=11, value=budget_display)
            ws.cell(row=row_num, column=12, value=urgency_display)
            ws.cell(row=row_num, column=13, value=hotness_display)
            ws.cell(row=row_num, column=14, value=tier_display)
            ws.cell(row=row_num, column=15, value=budget_signal_text)
            ws.cell(row=row_num, column=16, value=urgency_signal_text)
            serial += 1
            row_num += 1

    # Auto-fit columns
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_length + 4, 50)

    ws.auto_filter.ref = ws.dimensions

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


async def generate_csv(run_id: UUID, db: AsyncSession, scope: str = "all") -> BytesIO:
    companies = await _fetch_companies(run_id, db, scope=scope)

    output = StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "Serial#", "Company Name", "Website", "Geo/City",
        "Contact Name", "Designation", "LinkedIn", "Email",
        "Phone", "Final Score", "Budget Score", "Urgency Score",
        "Deal Hotness", "Hotness Tier", "Budget Signal", "Urgency Signal",
    ])

    serial = 1
    for company in companies:
        score_display = company.final_score if company.final_score is not None else "N/A"
        city_str = ", ".join(filter(None, [company.city, company.state_region, company.country]))
        budget_display = company.budget_signal_score if company.budget_signal_score is not None else ""
        urgency_display = company.urgency_signal_score if company.urgency_signal_score is not None else ""
        hotness_display = company.deal_hotness_score if company.deal_hotness_score is not None else ""
        tier_display = company.deal_hotness_tier or ""
        budget_signal_text = _extract_signal_text(company, "budget_signals")
        urgency_signal_text = _extract_signal_text(company, "urgency_signals")

        if company.contacts:
            for contact in company.contacts:
                writer.writerow([
                    serial, company.name, company.website, city_str,
                    contact.full_name, contact.designation, contact.linkedin_url,
                    contact.email, contact.phone, score_display,
                    budget_display, urgency_display, hotness_display, tier_display,
                    budget_signal_text, urgency_signal_text,
                ])
                serial += 1
        else:
            writer.writerow([
                serial, company.name, company.website, city_str,
                "", "", "", "", "", score_display,
                budget_display, urgency_display, hotness_display, tier_display,
                budget_signal_text, urgency_signal_text,
            ])
            serial += 1

    buffer = BytesIO(output.getvalue().encode("utf-8"))
    buffer.seek(0)
    return buffer
